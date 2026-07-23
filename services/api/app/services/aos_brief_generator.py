"""AOS brief generator — library FILL + verbatim PRESERVE (+ optional API enhancement).

Parser and Generator are SEPARATE. Default path uses paragraph library (no API).
Set AOD_AOS_USE_API=1 (and pass use_api=True) for optional Claude FILL enhancement.
"""

from __future__ import annotations

import logging
import os
import re
from datetime import date
from pathlib import Path
from typing import Any

from app.services.aos_paragraph_library import (
    build_certificate_of_service,
    evaluate_selection_logic,
    fill_section,
    library_loaded,
    load_paragraph_library,
)
from app.services.aos_section_prompts import build_section_prompt, fact_keys_included
from app.services.aos_system_prompt import SYSTEM_PROMPT_PART_8_1
from app.services.brief_parser.aos_discretionary import (
    BRIEF_TYPE,
    CANONICAL_PRESERVE_BY_SECTION,
    build_default_aos_template,
    is_empty_or_label_preserve,
    merge_canonical_preserve_sections,
)
from app.services.brief_parser.classification import slugify

LOGGER = logging.getLogger(__name__)

# Full Part 8.1 — never truncate when sending to Anthropic.
AOS_SYSTEM_PROMPT = SYSTEM_PROMPT_PART_8_1

LIBRARY_FILL_SECTION_IDS = frozenset(
    {
        "argument_intro",
        "section_a",
        "section_b",
        "section_d_adverse",
        "section_e_balancing",
        "certificate_of_service",
    }
)


def aos_use_api_env() -> bool:
    """Optional LLM FILL enhancement — default OFF (library path)."""
    return os.getenv("AOD_AOS_USE_API", "").strip().lower() in {"1", "true", "yes", "on"}


def flatten_facts(d: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    result: dict[str, Any] = {}
    for k, v in (d or {}).items():
        full_key = f"{prefix}{k}" if not prefix else f"{prefix}_{k}"
        if isinstance(v, dict):
            result.update(flatten_facts(v, full_key))
        else:
            result[full_key] = v
            # Also expose leaf key for ##APPLICANT_FULL_NAME## style tokens
            result[str(k)] = v
    return result


def apply_simple_substitutions(template_text: str, facts: dict[str, Any]) -> tuple[str, list[str]]:
    flat = flatten_facts(facts.get("case_facts", facts) if isinstance(facts, dict) else {})
    # Normalize keys to lowercase for lookup
    flat_lower = {str(k).lower(): v for k, v in flat.items()}
    # Common aliases from practice-area fact field ids
    aliases = {
        "applicant_full_name": flat_lower.get("applicant_full_name")
        or flat_lower.get("full_name")
        or flat_lower.get("applicantfullname"),
        "entry_date": flat_lower.get("entry_date") or flat_lower.get("last_entry_date") or flat_lower.get("entrydate"),
        "port_of_entry": flat_lower.get("port_of_entry") or flat_lower.get("last_entry_port") or flat_lower.get("portofentry"),
        "entry_visa_type": flat_lower.get("entry_visa_type")
        or flat_lower.get("last_entry_visa_type")
        or flat_lower.get("entryvisatype"),
        "petitioner_name": flat_lower.get("petitioner_name") or flat_lower.get("qualifyingrelative"),
        "petitioner_relationship": flat_lower.get("petitioner_relationship"),
        "i130_approved_date": flat_lower.get("i130_approved_date") or flat_lower.get("i130approveddate"),
        "i485_filed_date": flat_lower.get("i485_filed_date") or flat_lower.get("i485fileddate"),
        "case_theme": flat_lower.get("case_theme") or flat_lower.get("casetheme"),
        "case_theme_brief": flat_lower.get("case_theme_brief") or flat_lower.get("casethemebrief"),
        "adverse_factor_brief": flat_lower.get("adverse_factor_brief") or flat_lower.get("adversefactorbrief"),
        "section_a_heading": flat_lower.get("section_a_heading") or flat_lower.get("sectionaheading"),
        "section_a_facts": flat_lower.get("section_a_facts") or flat_lower.get("sectionafacts"),
        "section_b_heading": flat_lower.get("section_b_heading") or flat_lower.get("sectionbheading"),
        "section_b_facts": flat_lower.get("section_b_facts") or flat_lower.get("sectionbfacts"),
        "adverse_heading": flat_lower.get("adverse_heading") or flat_lower.get("adverseheading"),
        "adverse_facts": flat_lower.get("adverse_facts")
        or flat_lower.get("adversefacts")
        or flat_lower.get("adversfactors"),
        "balancing_inventory": flat_lower.get("balancing_inventory")
        or flat_lower.get("balancinginventory")
        or flat_lower.get("positiveequities"),
        "departure_harm": flat_lower.get("departure_harm") or flat_lower.get("departureharm"),
        "departure_bar_type": flat_lower.get("departure_bar_type")
        or flat_lower.get("departurebartype")
        or "three-year or ten-year",
        "applicant_pronoun_subject": flat_lower.get("applicant_pronoun_subject")
        or flat_lower.get("pronoun_subject")
        or flat_lower.get("pronounsubject")
        or "the applicant",
        "applicant_pronoun_object": flat_lower.get("applicant_pronoun_object")
        or flat_lower.get("pronoun_object")
        or "the applicant",
        "applicant_pronoun_possessive": flat_lower.get("applicant_pronoun_possessive")
        or flat_lower.get("pronoun_possessive")
        or "the applicant's",
        "bar_number": flat_lower.get("bar_number")
        or flat_lower.get("attorney_bar")
        or flat_lower.get("attorneybar"),
        "parties_served": flat_lower.get("parties_served") or flat_lower.get("partiesserved") or "USCIS",
        "service_method": flat_lower.get("service_method")
        or flat_lower.get("servicemethod")
        or "electronic submission",
        "attorney_name": flat_lower.get("attorney_name") or flat_lower.get("attorneyname"),
        "firm_name": flat_lower.get("firm_name") or flat_lower.get("firmname"),
        "attorney_bar": flat_lower.get("attorney_bar") or flat_lower.get("bar_number") or flat_lower.get("attorneybar"),
        "date": flat_lower.get("date") or flat_lower.get("prepared_date") or date.today().isoformat(),
        "applicant_a_number": flat_lower.get("applicant_a_number") or flat_lower.get("a_number") or flat_lower.get("anumbers"),
    }
    for k, v in aliases.items():
        if v is not None and str(v).strip():
            flat_lower[k] = v

    result = template_text or ""
    unfilled: list[str] = []
    for match in re.finditer(r"##([A-Z0-9_]+)##", template_text or ""):
        key = match.group(1).lower()
        val = flat_lower.get(key)
        if val is not None and str(val).strip():
            result = result.replace(match.group(0), str(val))
        else:
            unfilled.append(match.group(1))
    return result, unfilled


def get_fact(facts: dict[str, Any], key: str) -> str:
    flat = flatten_facts(facts.get("case_facts", facts) if isinstance(facts, dict) else {})
    flat_lower = {str(k).lower(): v for k, v in flat.items()}
    # camelCase field ids from practice-area facts
    camel = "".join(p.capitalize() if i else p for i, p in enumerate(key.split("_")))
    for candidate in (key, key.replace("_", ""), camel, key.lower()):
        val = flat_lower.get(candidate.lower())
        if val is not None and str(val).strip():
            return str(val).strip()
    # Direct fields map
    fields = facts.get("fields") if isinstance(facts, dict) else None
    if isinstance(fields, dict):
        for candidate in (key, camel, key.replace("_", "")):
            val = fields.get(candidate) or fields.get(camel)
            if val is not None and str(val).strip():
                return str(val).strip() if not isinstance(val, list) else "; ".join(str(x) for x in val)
        # snake → camel
        parts = key.split("_")
        camel_id = parts[0] + "".join(p.title() for p in parts[1:])
        val = fields.get(camel_id)
        if val is not None and str(val).strip():
            return str(val).strip() if not isinstance(val, list) else "; ".join(str(x) for x in val)
    return ""


def validate_required_inputs(facts: dict[str, Any], template: dict[str, Any]) -> list[str]:
    missing: list[str] = []
    for section in template.get("sections") or []:
        for slot in section.get("slots") or []:
            if not slot.get("required"):
                continue
            key = slot.get("replacement_key") or ""
            if not get_fact(facts, key):
                missing.append(key)
    return sorted(set(missing))


def assemble_preserve_section(section: dict[str, Any], facts: dict[str, Any]) -> str:
    """Always emit full preserved_text after slot substitution — never label placeholders."""
    sid = str(section.get("section_id") or "")
    text = section.get("preserved_text") or ""
    if is_empty_or_label_preserve(text):
        text = CANONICAL_PRESERVE_BY_SECTION.get(sid) or ""
    if not text:
        return ""
    text, _ = apply_simple_substitutions(text, facts)
    return text


def assemble_fill_section_no_api(section: dict[str, Any], facts: dict[str, Any]) -> str:
    """Default FILL path: paragraph library lookup (filing-quality prose, no API)."""
    sid = str(section.get("section_id") or "")

    if sid == "certificate_of_service" or "certificate" in sid:
        return build_certificate_of_service(facts)

    if sid in LIBRARY_FILL_SECTION_IDS and library_loaded():
        body = fill_section(sid, facts)
        if body and "[SECTION " not in body and "[FILL with matter facts]" not in body:
            return body

    if section.get("fill_template"):
        text, _ = apply_simple_substitutions(section["fill_template"], facts)
        return text
    if section.get("closing_template"):
        # Balancing with inventory + closing (fallback if library missing)
        if library_loaded():
            return fill_section("section_e_balancing", facts)
        inventory = get_fact(facts, "balancing_inventory") or get_fact(facts, "positive_equities")
        body_bits = []
        if inventory:
            body_bits.append(f"On the positive side of the ledger: {inventory}")
        close, _ = apply_simple_substitutions(section["closing_template"], facts)
        body_bits.append(close)
        return "\n\n".join(body_bits)

    if library_loaded() and sid:
        body = fill_section(sid, facts)
        if body and "[SECTION " not in body:
            return body

    required_slots = [s for s in section.get("slots") or [] if s.get("required")]
    slot_lines = []
    for s in required_slots:
        key = s.get("replacement_key") or ""
        label = s.get("label") or key
        slot_lines.append(f"  - {label}: {get_fact(facts, key) or '[FACT NEEDED]'}")
    heading = section.get("heading") or section.get("section_id") or "Section"
    for heading_key in ("section_a_heading", "section_b_heading", "adverse_heading"):
        if heading_key in {s.get("replacement_key") for s in required_slots}:
            authored = get_fact(facts, heading_key)
            if authored:
                heading = authored
    cites = ", ".join(section.get("citations_to_include") or [])
    return (
        f"[SECTION REQUIRES LIBRARY OR API PROSE]\n"
        f"Section: {heading}\n"
        f"Available facts:\n" + "\n".join(slot_lines) + "\n"
        f"Citations to include: {cites}\n"
        f"[END PLACEHOLDER]"
    )


def assemble_footnotes(template_footnotes: dict[str, Any], used_citations: list[str]) -> list[dict[str, Any]]:
    preserved = template_footnotes.get("preserved_citations") or []
    all_citations: dict[str, str] = {}
    for fn in preserved:
        if isinstance(fn, dict) and fn.get("key") and fn.get("text"):
            all_citations[str(fn["key"])] = str(fn["text"])
    for cite in used_citations:
        key = slugify(cite)
        if key and key not in all_citations:
            all_citations[key] = cite if cite.endswith(".") else f"{cite}."
    return [{"number": i, "key": k, "text": t} for i, (k, t) in enumerate(all_citations.items(), start=1)]


def validate_brief(
    sections_assembled: list[dict[str, Any]],
    facts: dict[str, Any],
    footnotes: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Part 9 quality validation — deterministic."""
    errors: list[str] = []
    warnings: list[str] = []

    case_theme = get_fact(facts, "case_theme")
    applicant_name = get_fact(facts, "applicant_full_name") or get_fact(facts, "applicantName")

    full_text = "\n\n".join(s.get("body", "") for s in sections_assembled)
    headings = [s.get("heading", "") for s in sections_assembled if s.get("heading")]

    balancing = next(
        (s for s in sections_assembled if "balancing" in (s.get("section_id") or "").lower()),
        None,
    )
    if balancing:
        if "This is not a case about" not in (balancing.get("body") or ""):
            errors.append("FAIL: Balancing section does not contain 'This is not a case about' closing.")
    else:
        errors.append("FAIL: Balancing section (III-E) not found in output.")

    conclusion = next(
        (s for s in sections_assembled if "conclusion" in (s.get("section_id") or "").lower()),
        None,
    )
    if conclusion and case_theme and case_theme[:30] not in (conclusion.get("body") or ""):
        warnings.append("WARN: Case theme may not appear in Conclusion section.")

    # Theme should appear ~4 times (cover/arg/balance/conclusion) — soft check
    if case_theme:
        theme_hits = full_text.count(case_theme[:40]) if len(case_theme) >= 40 else full_text.count(case_theme)
        if theme_hits < 2:
            warnings.append("WARN: Case theme appears fewer than twice in assembled body.")

    adverse = next(
        (
            s
            for s in sections_assembled
            if "adverse" in (s.get("section_id") or "").lower() or "section_d" in (s.get("section_id") or "").lower()
        ),
        None,
    )
    if adverse:
        h = (adverse.get("heading") or "").lower()
        for forbidden in ["immigration violation", "overstay", "unlawful presence"]:
            if forbidden in h:
                errors.append(f"FAIL: Adverse section heading contains forbidden word '{forbidden}'.")

    category_labels = [
        "family unity",
        "humanitarian concerns",
        "good moral character",
        "community ties",
        "employment history",
        "adverse factors",
        "immigration violations",
        "immigration history",
    ]
    for heading in headings:
        if (heading or "").lower().strip() in category_labels:
            errors.append(f"FAIL: Section heading '{heading}' is a category label, not an argument claim.")

    in_text_cite_pattern = r"\(\w+ of \w+,\s+\d+ I&N"
    for section in sections_assembled:
        body = section.get("body", "")
        if re.search(in_text_cite_pattern, body or ""):
            errors.append(
                f"FAIL: In-text citation found in section '{section.get('section_id')}'. "
                "Citations must be in footnotes only."
            )

    bracket_pattern = r"\[[^\]]{5,100}\]"
    for section in sections_assembled:
        for m in re.findall(bracket_pattern, section.get("body") or ""):
            if not any(x in m for x in ["INA", "8 U.S.C.", "§", "BIA", "I&N Dec.", "SECTION REQUIRES", "END PLACEHOLDER", "FACT NEEDED"]):
                warnings.append(f"WARN: Unresolved bracket in '{section.get('section_id')}': {m[:60]}...")

    if applicant_name and applicant_name not in full_text:
        errors.append(f"FAIL: Applicant name '{applicant_name}' not found in brief text.")

    for phrase, severity in [
        ("rebuttal", "ERROR"),
        ("in rebuttal", "ERROR"),
        ("unfortunately", "WARNING"),
        ("we regret", "WARNING"),
        ("while it is true that", "WARNING"),
        ("it should be noted that", "WARNING"),
        ("we wish to point out", "WARNING"),
    ]:
        if phrase in full_text.lower():
            if severity == "ERROR":
                errors.append(f"FAIL: Forbidden phrase '{phrase}' found in brief.")
            else:
                warnings.append(f"WARN: Weak phrasing '{phrase}' found — consider removing.")

    if footnotes is not None:
        fn_text = "\n".join(fn.get("text", "") for fn in footnotes)
        for required in ["Matter of Patel", "Matter of Marin", "Matter of Arai"]:
            if required not in fn_text and required not in full_text:
                warnings.append(f"WARN: Core citation '{required}' not found in footnotes.")

    return {
        "passed": len(errors) == 0,
        "error_count": len(errors),
        "warning_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
    }


def build_output_docx(
    sections_assembled: list[dict[str, Any]],
    footnotes: list[dict[str, Any]],
    facts: dict[str, Any],
    output_path: str | Path,
) -> str:
    from docx import Document  # type: ignore[import-not-found]
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]
    from docx.shared import Inches, Pt  # type: ignore[import-not-found]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    applicant = get_fact(facts, "applicant_full_name") or "Applicant"
    theme = get_fact(facts, "case_theme")
    a_number = get_fact(facts, "applicant_a_number")
    attorney = get_fact(facts, "attorney_name") or "Counsel"
    firm = get_fact(facts, "firm_name") or ""
    prepared = get_fact(facts, "date") or date.today().isoformat()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("MEMORANDUM IN SUPPORT OF APPLICATION FOR ADJUSTMENT OF STATUS")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Applicant: {applicant}")

    if a_number:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"File No. {a_number}")

    if theme:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(theme)
        run.italic = True
        run.font.size = Pt(11)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Submitted: {prepared}")
    p.add_run(f"\nPrepared by: {attorney}")
    if firm:
        p.add_run(f"\n{firm}")

    doc.add_page_break()

    for section_content in sections_assembled:
        heading_text = section_content.get("heading")
        body_text = section_content.get("body", "")
        level = int(section_content.get("level") or 1)
        if heading_text and heading_text != "__COVER__":
            try:
                doc.add_heading(heading_text, level=min(max(level, 1), 3))
            except Exception:
                hp = doc.add_paragraph(heading_text)
                if hp.runs:
                    hp.runs[0].bold = True
        if body_text:
            for para_text in str(body_text).split("\n\n"):
                if para_text.strip():
                    bp = doc.add_paragraph(para_text.strip())
                    bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_page_break()
    doc.add_heading("FOOTNOTES", level=1)
    for fn in footnotes:
        fp = doc.add_paragraph()
        run = fp.add_run(f"{fn['number']}. ")
        run.bold = True
        fp.add_run(fn["text"])
        fp.paragraph_format.left_indent = Inches(0.25)

    out = str(output_path)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def _infer_pronouns(fields: dict[str, Any], name: str) -> tuple[str, str, str]:
    """Best-effort pronouns from explicit fields or honorifics; default they/them/their."""
    subj = str(fields.get("pronounSubject") or fields.get("pronoun_subject") or "").strip().lower()
    obj = str(fields.get("pronounObject") or fields.get("pronoun_object") or "").strip().lower()
    poss = str(fields.get("pronounPossessive") or fields.get("pronoun_possessive") or "").strip().lower()
    if subj and obj and poss:
        return subj, obj, poss
    gender = str(fields.get("gender") or fields.get("applicantGender") or "").strip().upper()
    if gender in {"F", "FEMALE"}:
        return "she", "her", "her"
    if gender in {"M", "MALE"}:
        return "he", "him", "his"
    low = (name or "").lower()
    if low.startswith(("ms.", "mrs.", "miss ")):
        return "she", "her", "her"
    if low.startswith("mr."):
        return "he", "him", "his"
    return "they", "them", "their"


def _adverse_type(description: str) -> str:
    d = (description or "").lower()
    if "overstay" in d or "unlawful presence" in d or "overstay" in d:
        return "overstay"
    if "remov" in d or "deport" in d:
        return "prior_removal"
    if "convict" in d or "criminal" in d or "arrest" in d:
        return "criminal"
    return "other"


def _facts_from_drafting_fields(fields: dict[str, Any] | None, matter_id: str = "") -> dict[str, Any]:
    """Map practice-area fact field ids → full Part 6 case_facts shape."""
    f = fields or {}

    def g(*keys: str) -> str:
        for k in keys:
            v = f.get(k)
            if isinstance(v, list):
                if v:
                    return "; ".join(str(x) for x in v)
            elif v is not None and str(v).strip():
                return str(v).strip()
        return ""

    full_name = g("applicantName", "applicant_full_name")
    pronoun_s, pronoun_o, pronoun_p = _infer_pronouns(f, full_name)
    adverse_desc = g("adverseFacts", "adverseFactors", "adverse_facts")
    section_a_facts = g("sectionAFacts", "section_a_facts", "positiveEquities")
    section_b_facts = g("sectionBFacts", "section_b_facts")
    positive_inventory = g("balancingInventory", "positiveEquities", "balancing_inventory")
    hardship = g("extremeHardshipFactors", "hardship", "humanitarianFactors")
    petitioner = g("petitionerName", "qualifyingRelative", "petitioner_name")
    relationship = g("petitionerRelationship", "petitioner_relationship")

    # Seed family_ties from petitioner + narrative equities (Part 6 shape).
    family_members: list[dict[str, Any]] = []
    if petitioner:
        quality = section_a_facts[:500] if section_a_facts else ""
        family_members.append(
            {
                "name": petitioner,
                "relationship": relationship or "qualifying relative",
                "status": (
                    "U.S. citizen"
                    if "citizen" in (relationship or "").lower()
                    or "grandson" in (section_a_facts or "").lower()
                    or "granddaughter" in (section_a_facts or "").lower()
                    or "child" in (section_a_facts or "").lower()
                    else g("petitionerStatus") or "see petition facts"
                ),
                "dependence": g("familyDependence") or "see Section A facts",
                "quality_description": quality,
            }
        )
    # If narrative mentions autistic dependent but petitioner is caregiver relative,
    # ensure autism flag is detectable for library selection.
    if section_a_facts and ("autism" in section_a_facts.lower() or "autistic" in section_a_facts.lower()):
        if not any("autism" in str(m.get("quality_description") or "").lower() for m in family_members):
            family_members.append(
                {
                    "name": g("dependentName") or "U.S. citizen dependent",
                    "relationship": g("dependentRelationship") or "grandson",
                    "status": "US_citizen",
                    "dependence": "primary caregiver for autistic dependent",
                    "quality_description": section_a_facts[:500],
                }
            )
    family_notes = g("familyTies", "family_ties", "positiveEquities")

    para_sels = f.get("paragraphSelections") or f.get("paragraph_selections") or {}
    if not isinstance(para_sels, dict):
        para_sels = {}

    return {
        "case_facts": {
            "meta": {
                "matter_id": matter_id,
                "prepared_by": g("attorneyName", "attorney_name"),
                "prepared_date": date.today().isoformat(),
                "brief_type": BRIEF_TYPE,
            },
            "applicant": {
                "full_name": full_name,
                "a_number": g("aNumber", "applicant_a_number"),
                "date_of_birth": g("applicantDob", "date_of_birth"),
                "country_of_birth": g("countryOfBirth"),
                "country_of_citizenship": g("countryOfCitizenship"),
                "gender": g("gender", "applicantGender") or "X",
                "pronoun_subject": pronoun_s,
                "pronoun_object": pronoun_o,
                "pronoun_possessive": pronoun_p,
            },
            "entry_and_immigration_history": {
                "last_entry_date": g("entryDate", "entry_date"),
                "last_entry_port": g("portOfEntry", "port_of_entry"),
                "last_entry_visa_type": g("entryVisaType", "entry_visa_type"),
                "authorized_stay_expiration": g("authorizedStayExpiration"),
                "overstay_start_date": g("overstayStartDate"),
                "prior_removal_orders": bool(g("priorRemovalOrders")),
                "prior_immigration_violations": g("priorFilings", "priorImmigrationViolations") or None,
                "departure_would_trigger_bar": True,
            },
            "petition": {
                "petitioner_name": petitioner,
                "petitioner_us_citizen_date": g("petitionerUscDate", "petitionerUsCitizenDate"),
                "petitioner_relationship": relationship,
                "i130_approved_date": g("i130ApprovedDate", "i130_approved_date"),
                "i485_filed_date": g("i485FiledDate", "i485_filed_date"),
                "i130_filed_date": g("i130FiledDate"),
                "i130_receipt_number": g("i130ReceiptNumber"),
                "i485_receipt_number": g("i485ReceiptNumber"),
            },
            "attorney": {
                "attorney_name": g("attorneyName", "attorney_name"),
                "firm_name": g("firmName", "firm_name"),
                "bar_number": g("attorneyBar", "barNumber", "attorney_bar"),
                "email": g("attorneyEmail"),
                "phone": g("attorneyPhone"),
                "address": g("attorneyAddress"),
            },
            "case_architecture": {
                "case_theme": g("caseTheme", "case_theme"),
                "case_theme_brief": g("caseThemeBrief", "case_theme_brief") or g("caseTheme"),
                "adverse_factor_brief": g("adverseFactorBrief", "adverse_factor_brief"),
                "section_a_heading": g("sectionAHeading", "section_a_heading"),
                "section_a_facts": section_a_facts,
                "section_b_heading": g("sectionBHeading", "section_b_heading"),
                "section_b_facts": section_b_facts,
                "adverse_heading": g("adverseHeading", "adverse_heading"),
                "balancing_inventory": positive_inventory,
                "include_aos_mechanism": True,
                "include_edwards": "criminal" in adverse_desc.lower(),
                "include_mendez": True,
            },
            "positive_factors": {
                "family_ties": {
                    "members": family_members,
                    "quality_notes": family_notes or section_a_facts,
                },
                "humanitarian": {
                    "hardship_if_denied": hardship,
                    "narrative": hardship,
                    "health_conditions": [hardship] if hardship else [],
                },
                "employment_and_economic": {
                    "narrative": section_b_facts or positive_inventory,
                    "us_employment_history": [section_b_facts] if section_b_facts else [],
                    "tax_history": g("taxHistory"),
                },
                "community_and_moral_character": {
                    "narrative": section_b_facts or positive_inventory,
                    "service_activities": [section_b_facts] if section_b_facts else [],
                    "criminal_record": "criminal" in adverse_desc.lower(),
                    "criminal_record_details": adverse_desc if "criminal" in adverse_desc.lower() else "",
                },
            },
            "adverse_factors": {
                "primary_adverse": {
                    "type": _adverse_type(adverse_desc),
                    "description": adverse_desc,
                    "context": adverse_desc,
                    "date_arose": g("adverseDate"),
                    "is_fraud": "fraud" in adverse_desc.lower() or "misrepresent" in adverse_desc.lower(),
                    "rehabilitation": g("rehabilitation"),
                },
                "additional_adverse": (
                    [{"type": "other", "description": g("adverseFactors"), "context": g("adverseFactors")}]
                    if g("adverseFactors") and g("adverseFactors") != adverse_desc
                    else []
                ),
            },
            "departure_harm": g("departureHarm", "departure_harm"),
            "departure_bar_type": g("departureBarType", "departure_bar_type") or "three-year or ten-year",
            "evidence_index": [],
            "paragraph_selections": para_sels,
            # Flat aliases for ##TOKEN## substitution
            "applicant_full_name": full_name,
            "applicant_a_number": g("aNumber"),
            "applicant_pronoun_subject": pronoun_s,
            "applicant_pronoun_object": pronoun_o,
            "applicant_pronoun_possessive": pronoun_p,
            "entry_date": g("entryDate"),
            "port_of_entry": g("portOfEntry"),
            "entry_visa_type": g("entryVisaType"),
            "petitioner_name": petitioner,
            "petitioner_relationship": relationship,
            "i130_approved_date": g("i130ApprovedDate"),
            "i485_filed_date": g("i485FiledDate"),
            "case_theme": g("caseTheme"),
            "case_theme_brief": g("caseThemeBrief", "caseTheme"),
            "adverse_factor_brief": g("adverseFactorBrief"),
            "section_a_heading": g("sectionAHeading"),
            "section_a_facts": section_a_facts,
            "section_b_heading": g("sectionBHeading"),
            "section_b_facts": section_b_facts,
            "adverse_heading": g("adverseHeading"),
            "adverse_facts": adverse_desc,
            "balancing_inventory": positive_inventory,
            "departure_harm": g("departureHarm"),
            "attorney_name": g("attorneyName"),
            "firm_name": g("firmName"),
            "attorney_bar": g("attorneyBar", "barNumber"),
            "bar_number": g("attorneyBar", "barNumber"),
            "date": date.today().isoformat(),
            # Preserve every raw drafting field under raw_fields for appendix completeness
            "raw_drafting_fields": {k: v for k, v in f.items() if v not in (None, "", [])},
        },
        "fields": f,
        "paragraph_selections": para_sels,
    }


def generate_aos_brief(
    template: dict[str, Any] | None,
    client_facts: dict[str, Any],
    output_path: str | Path,
    *,
    use_api: bool = False,
) -> dict[str, Any]:
    """
    Generate AOS discretionary brief.

    Default path: PRESERVE verbatim + FILL via paragraph library (no API).
    use_api: when True AND AOD_AOS_USE_API / caller intent AND Anthropic configured,
    optionally enhance FILL sections with Part 8.1/8.2 + thinking. Library remains
    the fallback if API fails.
    """
    tmpl = merge_canonical_preserve_sections(template or build_default_aos_template())
    # Accept drafting fields payload
    if "case_facts" not in client_facts and "fields" in client_facts:
        facts = _facts_from_drafting_fields(
            client_facts.get("fields"),  # type: ignore[arg-type]
            matter_id=str(client_facts.get("matter_id") or ""),
        )
        # Merge attorney from top-level if present
        for k in ("attorney_name", "firm_name", "attorney_bar"):
            if client_facts.get(k):
                facts["case_facts"][k] = client_facts[k]
                facts["case_facts"]["attorney"][k if k != "attorney_bar" else "bar_number"] = client_facts[k]
    else:
        facts = client_facts

    # Attorney selection-menu overrides
    if isinstance(client_facts, dict) and client_facts.get("paragraph_selections"):
        facts.setdefault("case_facts", facts if "applicant" in facts else facts.get("case_facts", {}))
        if isinstance(facts.get("case_facts"), dict):
            facts["case_facts"]["paragraph_selections"] = client_facts["paragraph_selections"]
        facts["paragraph_selections"] = client_facts["paragraph_selections"]

    # Defaults for PRESERVE slots
    cf = facts.get("case_facts") if isinstance(facts.get("case_facts"), dict) else None
    if isinstance(cf, dict):
        cf.setdefault("departure_bar_type", "three-year or ten-year")
        if not cf.get("bar_number") and cf.get("attorney_bar"):
            cf["bar_number"] = cf["attorney_bar"]
        app = cf.get("applicant") if isinstance(cf.get("applicant"), dict) else {}
        if app:
            cf.setdefault("applicant_pronoun_subject", app.get("pronoun_subject") or "the applicant")

    missing = validate_required_inputs(facts, tmpl)
    # Soft: do not hard-fail generation — placeholders remain for missing fields

    from app.services.llm import (
        aos_max_tokens,
        aos_model_name,
        aos_thinking_kwargs,
        generate_text,
        is_configured as llm_configured,
    )

    # Library is default; API only when caller requests AND AOD_AOS_USE_API=1 AND key set.
    api_available = bool(use_api and aos_use_api_env() and llm_configured())
    thinking_cfg = aos_thinking_kwargs() if api_available else None
    model_used = aos_model_name() if api_available else None
    fact_keys = fact_keys_included(facts)
    selection_report = evaluate_selection_logic(facts, load_paragraph_library())

    assembled: list[dict[str, Any]] = []
    used_cites: list[str] = []
    api_sections: list[str] = []
    api_errors: list[str] = []
    library_sections: list[str] = []

    for section in tmpl.get("sections") or []:
        classification = section.get("classification")
        body = ""
        heading = section.get("heading")

        # Attorney-authored headings override placeholders
        sid = section.get("section_id") or ""
        if sid == "section_a" and get_fact(facts, "section_a_heading"):
            heading = get_fact(facts, "section_a_heading")
        elif sid == "section_b" and get_fact(facts, "section_b_heading"):
            heading = get_fact(facts, "section_b_heading")
        elif sid == "section_d_adverse" and get_fact(facts, "adverse_heading"):
            heading = get_fact(facts, "adverse_heading")

        if classification == "PRESERVE":
            body = assemble_preserve_section(section, facts)
        elif classification == "CAPTION":
            body, _ = apply_simple_substitutions(section.get("fill_template") or "", facts)
        elif classification == "BOILERPLATE":
            if sid == "certificate_of_service" or (not section.get("preserved_text") and "certificate" in sid):
                body = build_certificate_of_service(facts)
                library_sections.append(sid)
            elif is_empty_or_label_preserve(section.get("preserved_text")) and sid in CANONICAL_PRESERVE_BY_SECTION:
                body = assemble_preserve_section(section, facts)
            elif section.get("preserved_text"):
                body, _ = apply_simple_substitutions(section["preserved_text"], facts)
            else:
                body = assemble_fill_section_no_api(section, facts)
        elif classification == "FILL":
            if section.get("api_assistance") == "required" and api_available:
                try:
                    user_prompt = build_section_prompt(section, facts)
                    raw = generate_text(
                        system=AOS_SYSTEM_PROMPT,
                        user=user_prompt,
                        max_tokens=aos_max_tokens(),
                        model=model_used,
                        thinking=thinking_cfg,
                    )
                    if not raw:
                        raise RuntimeError("empty Anthropic response")
                    parsed = _parse_api_response(raw)
                    body = parsed.get("body") or assemble_fill_section_no_api(section, facts)
                    if parsed.get("heading"):
                        heading = parsed["heading"]
                    used_cites.extend(parsed.get("footnotes") or [])
                    api_sections.append(sid)
                except Exception as exc:
                    LOGGER.warning("AOS FILL API failed for %s: %s", sid, exc)
                    api_errors.append(f"{sid}: {exc}")
                    body = assemble_fill_section_no_api(section, facts)
                    library_sections.append(sid)
            else:
                body = assemble_fill_section_no_api(section, facts)
                if sid in LIBRARY_FILL_SECTION_IDS:
                    library_sections.append(sid)
        else:
            body = assemble_preserve_section(section, facts) if section.get("preserved_text") else ""

        for cite in section.get("citations_to_include") or []:
            used_cites.append(cite)
        for cite in section.get("citations_embedded") or []:
            used_cites.append(cite)

        assembled.append(
            {
                "section_id": sid,
                "heading": heading,
                "level": section.get("level", 1),
                "body": body,
                "classification": classification,
            }
        )

    footnotes = assemble_footnotes(tmpl.get("footnotes") or {}, used_cites)
    path = build_output_docx(assembled, footnotes, facts, output_path)
    validation = validate_brief(assembled, facts, footnotes)

    return {
        "output_path": path,
        "validation": validation,
        "api_used": api_available and bool(api_sections),
        "api_requested": api_available,
        "api_sections": api_sections,
        "api_errors": api_errors,
        "library_sections": library_sections,
        "library_loaded": library_loaded(),
        "selection_logic": selection_report,
        "sections_generated": len(assembled),
        "footnotes_count": len(footnotes),
        "missing_fields": missing,
        "assembled_sections": assembled,
        "brief_type": BRIEF_TYPE,
        "prompt_meta": {
            "system_prompt_chars": len(AOS_SYSTEM_PROMPT),
            "system_prompt_source": "Part 8.1 SYSTEM_PROMPT_PART_8_1",
            "thinking": thinking_cfg,
            "model": model_used,
            "max_tokens": aos_max_tokens() if api_available else None,
            "fact_keys_included": fact_keys,
            "fact_key_count": len(fact_keys),
            "fill_path": "api" if api_sections else "paragraph_library",
            "aos_use_api_env": aos_use_api_env(),
        },
    }


def _build_section_user_prompt(section: dict[str, Any], facts: dict[str, Any]) -> str:
    """Backward-compatible alias — Part 8.2 build_section_prompt."""
    return build_section_prompt(section, facts)


def _parse_api_response(response_text: str) -> dict[str, Any]:
    result: dict[str, Any] = {"heading": None, "body": "", "footnotes": []}
    lines = (response_text or "").strip().split("\n")
    current = None
    buffer: list[str] = []
    for line in lines:
        if line.startswith("HEADING:"):
            current = "heading"
            result["heading"] = line[8:].strip()
        elif line.startswith("BODY:"):
            current = "body"
            remaining = line[5:].strip()
            buffer = [remaining] if remaining else []
        elif line.startswith("FOOTNOTES:"):
            if current == "body":
                result["body"] = "\n".join(buffer).strip()
                buffer = []
            current = "footnotes"
        else:
            if current in ("body", "footnotes"):
                buffer.append(line)
    if current == "body" and buffer:
        result["body"] = "\n".join(buffer).strip()
    elif current == "footnotes" and buffer:
        result["footnotes"] = [l.strip() for l in buffer if l.strip()]
    return result


def format_aos_writing_rules_for_prompt() -> str:
    """Inject into drafting agent for AOS SKU — full Part 8.1, never truncated."""
    return (
        "## AOS Brief Writing Rules (System Guide — PRESERVE vs FILL)\n"
        "- PRESERVE: Legal Standard (Patel/Marin/Arai + 1 USCIS-PM E.8) — copy/adapt only for firm voice; do not invent new law.\n"
        "- FILL: Statutory eligibility, Argument equities A/B, adverse D, balancing E — use matter facts.\n"
        "- Case theme appears FOUR times: cover subtitle, Argument opening, balancing close, Conclusion.\n"
        "- Adverse heading MUST NOT contain 'Immigration Violations', 'Overstay', or 'Unlawful Presence'.\n"
        "- Citations in footnotes only — no in-text parentheticals.\n"
        "- Never use 'rebuttal' or apologetic phrasing.\n"
        "- Balancing MUST close: 'This is not a case about [adverse]. It is a case about [theme]. "
        "A favorable exercise of discretion is both legally supported and compelled by the facts of this record.'\n"
        "\n## Full Part 8.1 System Prompt (verbatim)\n"
        f"{AOS_SYSTEM_PROMPT}"
    )


def assembled_sections_to_memo(assembled: list[dict[str, Any]]) -> str:
    """Flatten generator sections into a memo string for agent summary / UI."""
    parts: list[str] = []
    for s in assembled:
        heading = s.get("heading") or ""
        body = (s.get("body") or "").strip()
        if heading and heading != "__COVER__":
            parts.append(f"{heading}\n\n{body}" if body else str(heading))
        elif body:
            parts.append(body)
    return "\n\n".join(parts).strip()

