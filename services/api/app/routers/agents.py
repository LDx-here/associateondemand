"""Multi-agent endpoints — PM orchestrator, research, pattern, strategy, jobs."""

from __future__ import annotations

import os
from typing import Any, Literal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.agents._context import load_constitution
from app.agents.pattern_agent import build_knowledge_graph, run_pattern, seed_all, seed_from_dev_data
from app.agents.pm_orchestrator import dispatch, get_job_result, list_inbox, process_next_job
from app.agents.research_agent import run_research
from app.agents.strategy_agent import run_strategy
from app.services.document_linter import lint_document
from app.db import get_db
from app.models.agent_result import AgentResult
from app.services.redis_queue import get_job, queue_depth

router = APIRouter(prefix="/agents", tags=["agents"])


class PmDispatchRequest(BaseModel):
    matter_id: str
    instruction: str
    priority: str = "normal"
    async_mode: bool = False


class ResearchRequest(BaseModel):
    matter_id: str
    query: str
    sources: list[str] = Field(default_factory=lambda: ["constitution", "obsidian", "airtable"])


@router.get("/constitution/preview")
def constitution_preview(max_chars: int = 2000) -> dict[str, str]:
    text = load_constitution(max_chars_per_file=max_chars)
    return {"preview": text[:max_chars], "truncated": len(text) > max_chars}


@router.post("/pm/dispatch", response_model=AgentResult)
def pm_dispatch(body: PmDispatchRequest, db: Session = Depends(get_db)) -> AgentResult:
    return dispatch(
        db,
        matter_id=body.matter_id,
        instruction=body.instruction,
        priority=body.priority,
        run_sync=not body.async_mode,
    )


@router.get("/pm/inbox")
def pm_inbox(db: Session = Depends(get_db)) -> dict[str, Any]:
    return {"items": list_inbox(db), "queue_depth": queue_depth()}


@router.post("/jobs/process", response_model=AgentResult | None)
def process_job(db: Session = Depends(get_db)) -> AgentResult | None:
    return process_next_job(db)


@router.get("/jobs/{job_id}")
def job_status(job_id: str, db: Session = Depends(get_db)) -> dict[str, Any]:
    redis_job = get_job(job_id)
    if not redis_job:
        raise HTTPException(status_code=404, detail="Job not found")
    return {"redis": redis_job, "result": get_job_result(job_id)}


@router.post("/research/run", response_model=AgentResult)
def research_run(body: ResearchRequest) -> AgentResult:
    return run_research(body.matter_id, body.query, body.sources)


@router.post("/pattern/analyze", response_model=AgentResult)
def pattern_analyze(matter_id: str, facts: str = "") -> AgentResult:
    return run_pattern(matter_id, facts=facts or None)


@router.post("/pattern/seed")
def pattern_seed(source: str = "all") -> dict[str, int]:
    if source == "dev":
        return {"indexed": seed_from_dev_data(), "source": "dev"}
    result = seed_all()
    result["source"] = "all"
    return result


class MemoExportBody(BaseModel):
    matter_id: str = ""
    memo_text: str = ""
    format: Literal["docx", "txt"] = "docx"


def _memo_docx_blob(text: str) -> bytes:
    """Build a .docx from plain memo lines with Page X of Y footer (BUILD_SPEC §11)."""

    from io import BytesIO

    try:
        from docx import Document  # type: ignore[import-not-found]
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:
        raise RuntimeError("python-docx unavailable") from None

    def _add_field(paragraph, instruction: str) -> None:
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instruction
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)

    doc = Document()
    for line in text.replace("\r\n", "\n").split("\n"):
        doc.add_paragraph(line)

    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Page ")
    _add_field(paragraph, "PAGE")
    paragraph.add_run(" of ")
    _add_field(paragraph, "NUMPAGES")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.post("/research/memo-export")
def research_memo_export(body: MemoExportBody) -> Response:
    """Optional Word export for BUILD_SPEC §11 (falls back caller may use TXT)."""

    fname = (body.matter_id or "memo").replace("/", "_").replace(" ", "") + "_research_memo"

    memo = body.memo_text.strip()
    if not memo:
        raise HTTPException(status_code=400, detail="memo_text is required")

    lint_issues = lint_document(memo)
    if lint_issues:
        raise HTTPException(
            status_code=422,
            detail={"message": "Document linter failed", "issues": lint_issues},
        )

    if body.format == "txt":
        return Response(
            content=memo.encode("utf-8"),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{fname}.txt"'},
        )

    try:
        blob = _memo_docx_blob(memo)
    except RuntimeError:
        return Response(
            content=memo.encode("utf-8"),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{fname}.txt"'},
        )

    try:
        return Response(
            content=blob,
            media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            headers={"Content-Disposition": f'attachment; filename="{fname}.docx"'},
        )
    except Exception:
        return Response(
            content=memo.encode("utf-8"),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{fname}.txt"'},
        )


@router.get("/drafting/assessment-template")
def drafting_assessment_template() -> FileResponse:
    """Download the AOS Discretionary Factors case assessment workbook (4 tabs)."""

    from pathlib import Path

    candidates = [
        Path(os.getenv("AOD_ASSESSMENT_XLSX", "")),
        Path(__file__).resolve().parents[3] / "data" / "templates" / "AOS_Discretionary_Factors_Case_Assessment_Tool.xlsx",
        Path("/app/data/templates/AOS_Discretionary_Factors_Case_Assessment_Tool.xlsx"),
    ]
    for path in candidates:
        if path.is_file():
            return FileResponse(
                path,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename="AOS_Discretionary_Factors_Case_Assessment_Tool.xlsx",
            )
    raise HTTPException(status_code=404, detail="Assessment template not found — run scripts/generate-aos-assessment-template.py")


@router.post("/drafting/run", response_model=AgentResult)
def drafting_run(matter_id: str, instruction: str) -> AgentResult:
    from app.agents.drafting_agent import run_drafting

    return run_drafting(matter_id, instruction)


class CitationPackageBody(BaseModel):
    matter_id: str = ""
    memo_text: str = ""
    matter_label: str = ""


@router.post("/drafting/citation-package")
def drafting_citation_package(body: CitationPackageBody) -> dict[str, Any]:
    """Build citation verification manifest + reference PDFs from draft text."""

    memo = body.memo_text.strip()
    if not memo:
        raise HTTPException(status_code=400, detail="memo_text is required")

    from app.drafting.citation_extractor import resolve_sources_for_draft
    from app.drafting.citation_package import build_citation_package
    import tempfile
    import os

    sources = resolve_sources_for_draft(memo)
    out_dir = tempfile.mkdtemp(prefix="citation-pkg-", dir=os.getenv("UPLOAD_DIR", "/tmp/aod-uploads"))
    label = body.matter_label or body.matter_id or "Draft"
    pkg = build_citation_package(sources, out_dir, matter_label=label)
    if not pkg.get("ok"):
        raise HTTPException(status_code=503, detail=pkg.get("error", "Citation package build failed"))
    return pkg


class AosBriefExportBody(BaseModel):
    matter_id: str
    memo_text: str = ""
    client_name: str = ""
    a_number: str = ""
    case_theme: str = ""
    xlsx_path: str = ""


@router.post("/drafting/aos-brief-export")
def drafting_aos_brief_export(body: AosBriefExportBody) -> Response:
    """Export AOS discretionary brief as DOCX (from memo text or assessment workbook)."""

    from app.drafting.aos_brief_docx import build_aos_brief_docx, build_from_xlsx
    import tempfile
    import os

    fname = (body.matter_id or "aos_brief").replace("/", "_").replace(" ", "") + "_discretionary_brief.docx"

    if body.xlsx_path and os.path.isfile(body.xlsx_path):
        tmp = tempfile.mktemp(suffix=".docx")
        build_from_xlsx(body.xlsx_path, tmp, body.client_name or body.matter_id, body.a_number)
        with open(tmp, "rb") as fh:
            blob = fh.read()
        os.remove(tmp)
    else:
        memo = body.memo_text.strip()
        if not memo:
            raise HTTPException(status_code=400, detail="memo_text or xlsx_path is required")
        lint_issues = lint_document(memo)
        if lint_issues:
            raise HTTPException(
                status_code=422,
                detail={"message": "Document linter failed", "issues": lint_issues},
            )
        blob = build_aos_brief_docx(
            client_name=body.client_name or body.matter_id,
            matter_id=body.matter_id,
            draft_text=memo,
            case_theme=body.case_theme,
            a_number=body.a_number,
        )

    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.post("/strategy/recommend", response_model=AgentResult)
def strategy_recommend(matter_id: str, posture: str = "", selected_strategy: str = "") -> AgentResult:
    return run_strategy(matter_id, posture=posture, selected_strategy=selected_strategy or None)


@router.post("/strong-reader/run", response_model=AgentResult)
def strong_reader_run(
    matter_id: str,
    document_id: str,
    text: str,
    filename: str = "",
) -> AgentResult:
    from app.agents.strong_reader_agent import run_strong_reader

    return run_strong_reader(
        matter_id=matter_id,
        document_id=document_id,
        text=text,
        filename=filename,
    )


@router.get("/knowledge-map/graph")
def knowledge_map_graph() -> dict[str, Any]:
    return build_knowledge_graph()
