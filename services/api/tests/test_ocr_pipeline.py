"""OCR pipeline tests."""

from __future__ import annotations

import tempfile
from pathlib import Path

from app.pipelines.ocr_pipeline import run_ocr_pipeline, text_quality_score


def test_text_quality_empty() -> None:
    assert text_quality_score("") == 0.0


def test_plain_text_file() -> None:
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as fh:
        fh.write("This is a sufficiently long plain text document for intake testing purposes.")
        path = Path(fh.name)
    try:
        result = run_ocr_pipeline(path)
        assert result.method == "plain_text"
        assert "intake testing" in result.text
        assert result.processing_status == "processed"
    finally:
        path.unlink(missing_ok=True)


def test_docx_text_extraction() -> None:
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sample.docx"
        doc = Document()
        doc.add_paragraph("AOS Discretionary Brief Template")
        doc.add_paragraph("Respondent respectfully submits this memorandum under INA § 245(a).")
        doc.save(path)

        result = run_ocr_pipeline(path)
        assert result.method == "docx_text"
        assert result.processing_status == "processed"
        assert "AOS Discretionary Brief" in result.text
        assert "245(a)" in result.text


def test_legacy_doc_rejected_clearly() -> None:
    with tempfile.NamedTemporaryFile("wb", suffix=".doc", delete=False) as fh:
        fh.write(b"\xd0\xcf\x11\xe0 binary ole doc stub")
        path = Path(fh.name)
    try:
        result = run_ocr_pipeline(path)
        assert result.processing_status == "failed"
        assert "Legacy .doc" in str(result.metadata.get("reason", ""))
    finally:
        path.unlink(missing_ok=True)


def test_binary_zip_not_read_as_plaintext() -> None:
    """Office Open XML without .docx suffix must not return garbled UTF-8."""
    with tempfile.NamedTemporaryFile("wb", suffix=".bin", delete=False) as fh:
        fh.write(b"PK\x03\x04binary-zip-content-not-text")
        path = Path(fh.name)
    try:
        result = run_ocr_pipeline(path)
        assert result.processing_status == "failed"
        assert result.method == "unsupported"
        assert "Unsupported file type" in str(result.metadata.get("reason", ""))
    finally:
        path.unlink(missing_ok=True)
