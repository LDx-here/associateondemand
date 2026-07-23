"""Unit tests for AOS brief parser classification + generator validation."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

from app.services.aos_brief_generator import (
    assemble_fill_section_no_api,
    assemble_preserve_section,
    generate_aos_brief,
    validate_brief,
    validate_required_inputs,
)
from app.services.brief_parser import (
    BRIEF_TYPE,
    build_default_aos_template,
    classify_and_build_template,
    classify_paragraph,
    classify_section_by_heading,
    detect_variable_slots,
    extract_citations_from_paragraph,
    get_brief_type,
    group_blocks_into_sections,
    list_brief_types,
    resolve_section_classification,
)


def test_brief_type_registry_has_aos():
    spec = get_brief_type("AOS_DISCRETIONARY")
    assert spec is not None
    assert spec.catalog_sku == "aos-discretionary-brief"
    assert get_brief_type("aos-discretionary-brief") is spec
    assert any(s.brief_type == BRIEF_TYPE for s in list_brief_types())


def test_preserve_heading_legal_standard():
    assert classify_section_by_heading("I. LEGAL STANDARD") == "PRESERVE"
    assert classify_section_by_heading("A. Statutory Framework") == "PRESERVE"
    assert classify_section_by_heading("USCIS Policy Framework") == "PRESERVE"


def test_fill_heading_patterns():
    assert classify_section_by_heading("II. STATUTORY ELIGIBILITY") == "FILL"
    assert classify_section_by_heading("III. ARGUMENT") == "FILL"
    assert classify_section_by_heading("Family Ties Within the United States") == "FILL"
    assert classify_section_by_heading("E. The Balance of Equities") == "FILL"
    assert classify_section_by_heading("D. Adverse Considerations in Context") == "FILL"


def test_caption_and_boilerplate_headings():
    assert classify_section_by_heading("IN RE: Phanpit Sakkhi") == "CAPTION"
    assert classify_section_by_heading("Memorandum in Support of Application") == "CAPTION"
    assert classify_section_by_heading("IV. CONCLUSION") == "BOILERPLATE"
    assert classify_section_by_heading("Certificate of Service") == "BOILERPLATE"


def test_preserve_paragraph_bia_citation():
    text = (
        "The grant of an application for adjustment of status under section 245 is a matter "
        "of administrative grace. An applicant has the burden of showing that discretion "
        "should be exercised in his favor. Matter of Patel, 17 I&N Dec. 597 (BIA 1980)."
    )
    label, conf = classify_paragraph(text)
    assert label == "PRESERVE"
    assert conf > 0.5


def test_fill_paragraph_client_facts():
    text = (
        "Ms. Sakkhi has resided in the United States since her lawful admission on "
        "February 22, 2023. Her daughter, a U.S. citizen, filed an I-130 petition."
    )
    label, _ = classify_paragraph(text)
    assert label == "FILL"


def test_mixed_paragraph_is_fill_and_extracts_citations():
    text = (
        "As established in Matter of Marin, 16 I&N Dec. 581, 584 (BIA 1978), the factors "
        "to be weighed include family ties. Ms. Sakkhi's grandson Saeng, born September 2016, "
        "depends entirely on her care."
    )
    label, _ = classify_paragraph(text)
    assert label == "FILL"
    cites = extract_citations_from_paragraph(text)
    assert any("Marin" in c for c in cites)


def test_detect_variable_slots():
    text = (
        "Applicant A-123456789 entered on February 22, 2023 on a B-2 visa. "
        "[CASE THEME] applies."
    )
    slots = detect_variable_slots(text, known_applicant_name=None)
    keys = {s["replacement_key"] for s in slots}
    assert "applicant_a_number" in keys
    assert "entry_date" in keys
    assert "entry_visa_type" in keys
    assert "case_theme" in keys


def test_group_blocks_into_sections():
    blocks = [
        {"type": "paragraph", "text": "Cover line", "style": "Normal"},
        {"type": "heading1", "text": "I. LEGAL STANDARD", "style": "Heading 1"},
        {"type": "paragraph", "text": "Matter of Patel, 17 I&N Dec. 597 (BIA 1980).", "style": "Normal"},
        {"type": "heading2", "text": "II. STATUTORY ELIGIBILITY", "style": "Heading 2"},
        {"type": "paragraph", "text": "Ms. Sakkhi was admitted February 22, 2023.", "style": "Normal"},
    ]
    sections = group_blocks_into_sections(blocks)
    assert sections[0]["heading"] == "__COVER__"
    assert any(s["heading"] == "I. LEGAL STANDARD" for s in sections)
    assert resolve_section_classification("I. LEGAL STANDARD", sections[1]["paragraphs"]) == "PRESERVE"


def test_classify_and_build_template_schema():
    sections = [
        {
            "heading": "I. LEGAL STANDARD",
            "level": 1,
            "paragraphs": [
                "Adjustment is a matter of administrative grace. Matter of Patel, 17 I&N Dec. 597 (BIA 1980)."
            ],
        },
        {
            "heading": "II. STATUTORY ELIGIBILITY",
            "level": 1,
            "paragraphs": ["Ms. Sakkhi was admitted on February 22, 2023 on a B-2 visa."],
        },
    ]
    tmpl = classify_and_build_template(sections, source_name="test.docx")
    assert tmpl["brief_type"] == "AOS_DISCRETIONARY"
    assert tmpl["template_id"]
    assert len(tmpl["sections"]) >= 2
    assert tmpl["sections"][0]["classification"] == "PRESERVE"
    assert tmpl["sections"][1]["classification"] == "FILL"
    assert any(s.get("section_id") == "certificate_of_service" for s in tmpl["sections"])


def test_default_template_has_required_sections():
    tmpl = build_default_aos_template()
    ids = {s["section_id"] for s in tmpl["sections"]}
    assert "legal_standard" in ids
    assert "section_a" in ids
    assert "section_e_balancing" in ids
    assert "conclusion" in ids


def _sample_facts():
    return {
        "case_facts": {
            "applicant": {"full_name": "Phanpit Sakkhi", "a_number": "A-123456789"},
            "applicant_full_name": "Phanpit Sakkhi",
            "entry_date": "February 22, 2023",
            "port_of_entry": "Los Angeles",
            "entry_visa_type": "B-2",
            "petitioner_name": "Daughter USC",
            "petitioner_relationship": "daughter",
            "i130_approved_date": "August 1, 2023",
            "i485_filed_date": "September 1, 2023",
            "case_theme": (
                "This case concerns a retired nurse whose care of her autistic U.S. citizen "
                "grandson is indispensable."
            ),
            "case_theme_brief": "indispensable care of an autistic U.S. citizen grandson",
            "adverse_factor_brief": "an overstay",
            "section_a_heading": (
                "Ms. Sakkhi's Nursing Career Makes Her Uniquely Qualified to Care for "
                "Her Autistic U.S. Citizen Grandson"
            ),
            "section_a_facts": "She provides daily specialized care.",
            "section_b_heading": (
                "Ms. Sakkhi's Community Service and Moral Character Reinforce the Equities"
            ),
            "section_b_facts": "Temple volunteer; tax compliance.",
            "adverse_heading": (
                "The Circumstances of Ms. Sakkhi's Continued Presence Do Not Diminish "
                "the Strength of This Application"
            ),
            "adverse_facts": "She remained beyond authorized stay after changed circumstances.",
            "balancing_inventory": "family care, community service, clean record",
            "departure_harm": "Departure would trigger the 3/10-year bar and sever caregiving.",
            "attorney_name": "La'Dajia Ferguson",
            "firm_name": "Kingdom Counsel Firm",
            "attorney_bar": "TX 123",
            "date": "2026-07-22",
            "case_architecture": {
                "case_theme": (
                    "This case concerns a retired nurse whose care of her autistic U.S. citizen "
                    "grandson is indispensable."
                ),
                "adverse_factor_brief": "an overstay",
                "case_theme_brief": "indispensable care of an autistic U.S. citizen grandson",
            },
        }
    }


def test_assemble_preserve_legal_standard():
    tmpl = build_default_aos_template()
    legal = next(s for s in tmpl["sections"] if s["section_id"] == "legal_standard")
    body = assemble_preserve_section(legal, _sample_facts())
    assert "Matter of Patel" in body
    assert "Matter of Arai" in body


def test_assemble_fill_no_api_uses_template():
    tmpl = build_default_aos_template()
    elig = next(s for s in tmpl["sections"] if s["section_id"] == "statutory_eligibility")
    body = assemble_fill_section_no_api(elig, _sample_facts())
    assert "Phanpit Sakkhi" in body
    assert "February 22, 2023" in body


def test_validate_brief_forbidden_adverse_heading():
    sections = [
        {
            "section_id": "section_d_adverse",
            "heading": "Immigration Violations",
            "body": "She overstayed.",
        },
        {
            "section_id": "section_e_balancing",
            "heading": "Balancing",
            "body": "This is not a case about an overstay. It is a case about care.",
        },
        {
            "section_id": "conclusion",
            "heading": "Conclusion",
            "body": "Phanpit Sakkhi respectfully requests approval. This case concerns a retired nurse whose care of her autistic U.S. citizen grandson is indispensable.",
        },
    ]
    report = validate_brief(sections, _sample_facts())
    assert report["passed"] is False
    assert any("forbidden" in e.lower() or "Immigration" in e or "overstay" in e.lower() for e in report["errors"])


def test_validate_brief_rebuttal_forbidden():
    theme = _sample_facts()["case_facts"]["case_theme"]
    sections = [
        {
            "section_id": "section_e_balancing",
            "heading": "Balancing",
            "body": (
                "This is not a case about an overstay. It is a case about care. "
                "In rebuttal to this point, note also."
            ),
        },
        {
            "section_id": "conclusion",
            "heading": "IV. CONCLUSION",
            "body": f"Phanpit Sakkhi requests approval. {theme}",
        },
    ]
    report = validate_brief(sections, _sample_facts())
    assert any("rebuttal" in e.lower() for e in report["errors"])


def test_validate_brief_category_label_heading():
    sections = [
        {"section_id": "section_a", "heading": "Family Unity", "body": "x"},
        {
            "section_id": "section_e_balancing",
            "heading": "E",
            "body": "This is not a case about an overstay. It is a case about care.",
        },
    ]
    report = validate_brief(sections, _sample_facts())
    assert any("category label" in e.lower() for e in report["errors"])


def test_validate_brief_missing_balancing_close():
    sections = [
        {"section_id": "section_e_balancing", "heading": "E", "body": "Equities favor approval."},
    ]
    report = validate_brief(sections, _sample_facts())
    assert any("This is not a case about" in e for e in report["errors"])


def test_validate_brief_passes_happy_path():
    theme = _sample_facts()["case_facts"]["case_theme"]
    sections = [
        {
            "section_id": "argument_intro",
            "heading": "III. ARGUMENT",
            "body": f"Phanpit Sakkhi is eligible. {theme}",
        },
        {
            "section_id": "section_a",
            "heading": _sample_facts()["case_facts"]["section_a_heading"],
            "body": "Primary equity developed.",
        },
        {
            "section_id": "section_d_adverse",
            "heading": _sample_facts()["case_facts"]["adverse_heading"],
            "body": "She remained beyond authorized stay.",
        },
        {
            "section_id": "section_e_balancing",
            "heading": "E. Balancing",
            "body": (
                "This is not a case about an overstay. It is a case about "
                "indispensable care of an autistic U.S. citizen grandson. "
                "A favorable exercise of discretion is both legally supported and "
                "compelled by the facts of this record."
            ),
        },
        {
            "section_id": "conclusion",
            "heading": "IV. CONCLUSION",
            "body": f"For the foregoing reasons, Phanpit Sakkhi requests approval. {theme}",
        },
    ]
    report = validate_brief(sections, _sample_facts())
    assert report["passed"] is True, report["errors"]


def test_generate_aos_brief_without_api_writes_docx():
    facts = _sample_facts()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "brief.docx"
        result = generate_aos_brief(None, facts, out, use_api=False)
        assert Path(result["output_path"]).exists()
        assert result["sections_generated"] >= 5
        assert result["brief_type"] == "AOS_DISCRETIONARY"
        assert "validation" in result
        # Balancing close present even without API
        balancing = next(
            s for s in result["assembled_sections"] if s["section_id"] == "section_e_balancing"
        )
        assert "This is not a case about" in balancing["body"]
        legal = next(s for s in result["assembled_sections"] if s["section_id"] == "legal_standard")
        assert "Matter of Patel" in legal["body"]


def test_validate_required_inputs_reports_missing():
    tmpl = build_default_aos_template()
    missing = validate_required_inputs({"case_facts": {"applicant_full_name": "X"}}, tmpl)
    assert "case_theme" in missing or "entry_date" in missing


def test_parse_docx_roundtrip_if_python_docx():
    pytest.importorskip("docx")
    from docx import Document

    from app.services.brief_parser import parse_aos_brief_docx

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sample.docx"
        doc = Document()
        doc.add_heading("I. LEGAL STANDARD", level=1)
        doc.add_paragraph(
            "Adjustment is a matter of administrative grace. Matter of Patel, 17 I&N Dec. 597 (BIA 1980)."
        )
        doc.add_heading("II. STATUTORY ELIGIBILITY", level=1)
        doc.add_paragraph("Ms. Sakkhi was admitted on February 22, 2023 on a B-2 visa.")
        doc.save(path)
        tmpl = parse_aos_brief_docx(path)
        assert tmpl["brief_type"] == "AOS_DISCRETIONARY"
        classifications = [s["classification"] for s in tmpl["sections"]]
        assert "PRESERVE" in classifications
        assert "FILL" in classifications
