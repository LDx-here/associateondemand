# -*- coding: utf-8 -*-
"""
generate_aos_brief.py
Reads Tab 2 (Brief Development Matrix) and Tab 3 (Case Theme Worksheet) from
AOS_Discretionary_Factors_Case_Assessment_Tool.xlsx and produces a structured
AOS Discretionary Brief in .docx format.

USAGE:
    python3 generate_aos_brief.py [xlsx_path] [output_path] [client_name] [a_number]

HOW IT WORKS:
  Tab 2 (Brief Development Matrix) — issue spotting and weighing:
    • Column C  (Authority)          → citations for footnotes
    • Column D  (Legal Proposition)  → law for the Legal Standard section
    • Column H  (Weight)             → filter: Strong / Moderate / Limited rows only
    • Column J  (Draft Brief Para)   → argument text in the Analysis section
    • Column I  (Counterarg/Rebuttal)→ rebuttal text for adverse categories

  Tab 3 (Case Theme Worksheet) — narrative architecture:
    • Case Theme sentence (Part 1)   → used in brief header
    • Section A–D headings + mapped Tab 2 categories (Part 2) → primary argument sections
    • AOS Mechanism flag + heading (Part 3) → optional dedicated AOS section
    • Adverse heading + balancing notes (Part 4) → adverse + balancing section
    • Conclusion notes (Part 5)      → supplements boilerplate conclusion

FOOTNOTE SYSTEM:
  Superscript numbers in text; all citations collected in a FOOTNOTES section on
  a separate page at the end of the document. Never a References section.
  Duplicate citations deduplicate automatically (same authority = same footnote #).
"""

import sys
import re
import os
import zipfile
from collections import OrderedDict

import openpyxl
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Defaults ──────────────────────────────────────────────────────────────────
XLSX_PATH   = ("/sessions/fervent-zen-clarke/mnt/immigration/"
               "AOS Discretion (PM-602-0199)/"
               "AOS_Discretionary_Factors_Case_Assessment_Tool.xlsx")
OUTPUT_PATH = "/sessions/fervent-zen-clarke/mnt/outputs/AOS_Discretionary_Brief.docx"
CLIENT_NAME = "[APPLICANT NAME]"
A_NUMBER    = "[A-XXXXXXXXX]"

ACTIVE_WEIGHTS  = {"Strong", "Moderate", "Limited"}
SHEET_MATRIX    = "2. Brief Development Matrix"
SHEET_THEME     = "3. Case Theme"

# ── CLI overrides ─────────────────────────────────────────────────────────────
if len(sys.argv) > 1: XLSX_PATH   = sys.argv[1]
if len(sys.argv) > 2: OUTPUT_PATH = sys.argv[2]
if len(sys.argv) > 3: CLIENT_NAME = sys.argv[3]
if len(sys.argv) > 4: A_NUMBER    = sys.argv[4]


# ══════════════════════════════════════════════════════════════════════════════
# FOOTNOTE MANAGER
# ══════════════════════════════════════════════════════════════════════════════
class FootnoteManager:
    def __init__(self):
        self._notes = []
        self._seen  = {}

    def register(self, citation_text: str) -> int:
        key = citation_text.strip()
        if not key:
            return 0
        if key in self._seen:
            return self._seen[key]
        n = len(self._notes) + 1
        self._notes.append((n, key))
        self._seen[key] = n
        return n

    @property
    def notes(self):
        return list(self._notes)

FN = FootnoteManager()


# ══════════════════════════════════════════════════════════════════════════════
# READ TAB 2 — BRIEF DEVELOPMENT MATRIX
# ══════════════════════════════════════════════════════════════════════════════
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb[SHEET_MATRIX]

# Find header row
HEADER_ROW = None
for r in ws.iter_rows():
    for cell in r:
        if cell.value == "Category":
            HEADER_ROW = cell.row
            break
    if HEADER_ROW:
        break

if not HEADER_ROW:
    raise ValueError("Cannot find header row in Tab 2 — expected 'Category' in column A.")

header_map = {cell.value: cell.column for cell in ws[HEADER_ROW] if cell.value}

def col(name):
    return header_map.get(name)

COL_CAT   = col("Category")
COL_FACT  = col("Factor / Issue")
COL_AUTH  = col("Authority")
COL_PROP  = col("Legal Proposition")
COL_WT    = col("Weight")
COL_REB   = col("Government Counterargument / Our Rebuttal")
COL_DRAFT = col("Draft Brief Paragraph")

rows_data = []
for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
    category = r[COL_CAT  - 1] if COL_CAT  else None
    factor   = r[COL_FACT - 1] if COL_FACT else None
    if not factor or not category:
        continue
    rows_data.append({
        "category":  str(category).strip(),
        "factor":    str(factor).strip(),
        "authority": str(r[COL_AUTH  - 1]).strip() if COL_AUTH  and r[COL_AUTH  - 1] else "",
        "prop":      str(r[COL_PROP  - 1]).strip() if COL_PROP  and r[COL_PROP  - 1] else "",
        "weight":    str(r[COL_WT    - 1]).strip() if COL_WT    and r[COL_WT    - 1] else "",
        "rebuttal":  str(r[COL_REB   - 1]).strip() if COL_REB   and r[COL_REB   - 1] else "",
        "draft":     str(r[COL_DRAFT - 1]).strip() if COL_DRAFT and r[COL_DRAFT - 1] else "",
    })

# Group by category (order-preserving)
by_cat = OrderedDict()
for rd in rows_data:
    by_cat.setdefault(rd["category"], []).append(rd)

def is_adverse(cat_label: str) -> bool:
    low = cat_label.lower()
    return any(x in low for x in ("adverse", "violation", "fraud", "criminal",
                                   "national security", "public safety"))

def active_rows(cat_label: str) -> list:
    return [r for r in by_cat.get(cat_label, []) if r["weight"] in ACTIVE_WEIGHTS]

# All adverse categories that have active rows
all_adverse_cats = [c for c in by_cat if is_adverse(c) and active_rows(c)]


# ══════════════════════════════════════════════════════════════════════════════
# READ TAB 3 — CASE THEME WORKSHEET
# ══════════════════════════════════════════════════════════════════════════════
def _cv(sheet, row, col):
    """Return stripped cell value or empty string."""
    v = sheet.cell(row=row, column=col).value
    return str(v).strip() if v else ""

ct = wb[SHEET_THEME]

case_theme = _cv(ct, 5, 2)   # Part 1 — Case Theme sentence

# Part 2 — Primary sections (rows 9-12, cols B=heading, C=categories)
primary_sections = []
for r in range(9, 13):
    heading    = _cv(ct, r, 2)
    cats_raw   = _cv(ct, r, 3)
    if heading or cats_raw:
        cats = [c.strip() for c in cats_raw.split(",") if c.strip()]
        primary_sections.append({"heading": heading, "categories": cats})

# Part 3 — AOS mechanism
aos_include = _cv(ct, 15, 2).lower() not in ("no", "n", "false", "0")
aos_heading = _cv(ct, 16, 2) or "Adjustment of Status Is the Appropriate and Legally Available Relief"
aos_notes   = _cv(ct, 17, 2)

# Part 4 — Adverse / balancing
adv_heading   = _cv(ct, 20, 2) or "Any Negative Considerations Are Limited and Do Not Diminish the Applicant's Compelling Equities"
balancing_notes = _cv(ct, 21, 2)

# Part 5 — Conclusion
conclusion_notes = _cv(ct, 24, 2)


# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENT HELPERS
# ══════════════════════════════════════════════════════════════════════════════
doc = Document()

section = doc.sections[0]
section.page_width  = Inches(8.5)
section.page_height = Inches(11)
for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(section, attr, Inches(1.0))

def add_fn_ref(para, fn_num: int):
    if fn_num <= 0:
        return
    run = para.add_run(str(fn_num))
    run.font.superscript = True
    run.font.size = Pt(9)

def para(text="", bold=False, italic=False, size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_before=0, space_after=6, left_indent=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if left_indent:
        p.paragraph_format.left_indent = Inches(left_indent)
    if text:
        run = p.add_run(text)
        run.bold   = bold
        run.italic = italic
        run.font.size = Pt(size)
    return p

def heading(text, level=1, size=None, bold=True, color=None,
            space_before=12, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt({1: 13, 2: 12, 3: 11}.get(level, 11) if size is None else size)
    if color:
        run.font.color.rgb = color
    return p

def hrule():
    """Thin horizontal rule — pBdr inserted at correct schema position."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "BFBFBF")
    pBdr.append(bottom)
    _AFTER_PBDR = {qn(t) for t in ("w:shd", "w:tabs", "w:spacing", "w:ind",
                                    "w:jc", "w:outlineLvl", "w:rPr", "w:sectPr")}
    insert_before = next((child for child in pPr if child.tag in _AFTER_PBDR), None)
    if insert_before is not None:
        insert_before.addprevious(pBdr)
    else:
        pPr.append(pBdr)
    return p

def prop_para(prop_text: str, authority_text: str, left_indent=0.25):
    """Legal proposition with footnote superscript."""
    fn_num = FN.register(authority_text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Inches(left_indent)
    run = p.add_run(prop_text)
    run.font.size = Pt(12)
    if fn_num > 0:
        add_fn_ref(p, fn_num)
    return p

def analysis_para(draft_text: str, authority_text: str = "", left_indent=0.25):
    """Analysis paragraph with optional footnote."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Inches(left_indent)
    run = p.add_run(draft_text)
    run.font.size = Pt(12)
    if authority_text:
        fn_num = FN.register(authority_text)
        if fn_num > 0:
            add_fn_ref(p, fn_num)
    return p


# ══════════════════════════════════════════════════════════════════════════════
# COVER HEADER
# ══════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run("KINGDOM COUNSEL FIRM")
run.bold = True; run.font.size = Pt(11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
run = p.add_run("Immigration Law Practice")
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

hrule()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after  = Pt(2)
run = p.add_run("MEMORANDUM IN SUPPORT OF ADJUSTMENT OF STATUS")
run.bold = True; run.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run(f"In the Matter of {CLIENT_NAME}")
run.font.size = Pt(12); run.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run(f"File No.: {A_NUMBER}")
run.font.size = Pt(11)

if case_theme:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(case_theme)
    run.font.size = Pt(11); run.italic = True
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

hrule()


# ══════════════════════════════════════════════════════════════════════════════
# SECTION I — LEGAL STANDARD  (3 paragraphs — not a treatise)
# ══════════════════════════════════════════════════════════════════════════════
heading("I.  LEGAL STANDARD", level=1, space_before=16)

# Paragraph 1: Statutory eligibility
elig_auth = "INA §245(a) (8 U.S.C. §1255(a))"
prop_para(
    'Under INA §245(a), "[t]he status of an alien who was inspected and admitted or paroled into '
    'the United States . . . may be adjusted by the Attorney General, in his discretion," provided '
    'that: (1) the alien has applied; (2) the alien "is eligible to receive an immigrant visa and is '
    'admissible to the United States for permanent residence"; and (3) "an immigrant visa is '
    'immediately available . . . at the time his application is filed." INA §245(a).',
    elig_auth, left_indent=0
)

# Paragraph 2: Discretionary balancing standard
disc_auth = ("Matter of Patel, 17 I&N Dec. 597 (BIA 1980); "
             "Matter of Marin, 16 I&N Dec. 581, 584 (BIA 1978); "
             "Matter of Arai, 13 I&N Dec. 494, 496 (BIA 1970)")
prop_para(
    'Even where statutory eligibility is established, adjustment remains a discretionary benefit — '
    '"a matter of administrative grace where the applicant has the burden of showing that discretion '
    'should be exercised in his or her favor." Matter of Patel, 17 I&N Dec. 597 (BIA 1980). '
    'The analysis requires "[a] balancing of the negative factors evidencing the person\'s '
    'undesirability as a permanent resident with the social and humane considerations presented on '
    'his or her behalf to determine whether relief appears in the best interests of this country." '
    'Matter of Marin, 16 I&N Dec. 581, 584 (BIA 1978). "In the absence of adverse factors, '
    'adjustment will ordinarily be granted, still as a matter of discretion." '
    'Matter of Arai, 13 I&N Dec. 494, 496 (BIA 1970).',
    disc_auth, left_indent=0
)

# Paragraph 3: USCIS PM / totality
pm_auth = ("1 USCIS-PM E.8(A)-(C); USCIS Policy Memorandum PM-602-0199 (May 21, 2026); "
           "Matter of Marin, 16 I&N Dec. 581, 585 (BIA 1978)")
prop_para(
    'USCIS officers must weigh all positive and negative factors "in the totality of the record." '
    '1 USCIS-PM E.8(C)(3). "Meeting the statutory and regulatory requirements alone does not entitle '
    'the requestor to the benefit sought." 1 USCIS-PM E.8(A). Where negative factors are more '
    'serious, "unusual or even outstanding equities" may be required. '
    'Matter of Marin, 16 I&N Dec. 581, 585 (BIA 1978).',
    pm_auth, left_indent=0
)

hrule()


# ══════════════════════════════════════════════════════════════════════════════
# SECTION II — STATUTORY ELIGIBILITY (short)
# ══════════════════════════════════════════════════════════════════════════════
heading("II.  " + CLIENT_NAME + " IS STATUTORILY ELIGIBLE FOR ADJUSTMENT", level=1, space_before=16)

# Pull threshold/eligibility rows from Tab 2
thresh_rows = []
for cat, rows in by_cat.items():
    low = cat.lower()
    if "threshold" in low or "eligibility" in low or "why aos" in low:
        thresh_rows.extend([r for r in rows if r["weight"] in ACTIVE_WEIGHTS
                             and r["draft"] and not r["draft"].startswith("[This factor")])

if thresh_rows:
    for rd in thresh_rows:
        analysis_para(rd["draft"], rd["authority"], left_indent=0)
else:
    para(
        f"{CLIENT_NAME} satisfies all three statutory requirements for adjustment of status: "
        "(1) a timely Form I-485 has been filed; (2) an immigrant visa is immediately available "
        "based on an approved petition; and (3) the Applicant is admissible (or eligible for a "
        "waiver of any ground of inadmissibility). [Attorney to confirm each element.]",
        italic=True
    )

hrule()


# ══════════════════════════════════════════════════════════════════════════════
# SECTION III — ARGUMENT
# ══════════════════════════════════════════════════════════════════════════════
heading("III.  ARGUMENT", level=1, space_before=16)

# Opening thesis — with case theme if provided
intro_fn = FN.register("Matter of Marin, 16 I&N Dec. 581, 584–85 (BIA 1978); 1 USCIS-PM E.8(C)(2)")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(6)
p.add_run(
    f"A balanced review of the totality of the record establishes that {CLIENT_NAME} merits a "
    "favorable exercise of discretion. The Applicant presents compelling positive equities — "
    "detailed in the sections below — that, when weighed against any adverse factors in the record, "
    "support the conclusion that adjustment of status serves the best interests of the United States."
).font.size = Pt(12)
add_fn_ref(p, intro_fn)

# Thread the case theme as a standalone thematic statement
if case_theme:
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(10)
    p2.paragraph_format.left_indent  = Inches(0.25)
    run = p2.add_run(case_theme)
    run.font.size = Pt(12)
    run.italic    = True


# ── Primary Argument Sections (from Case Theme Worksheet) ─────────────────────
section_letters = "ABCDEFGH"
arg_idx = 0  # tracks which letter we're on

def write_argument_section(letter, section_heading, cat_list):
    """Write one primary argument section with analysis drawn from named Tab 2 categories."""
    heading(f"    {letter}.  {section_heading}", level=2, space_before=10, space_after=4)

    wrote_anything = False
    for cat_name in cat_list:
        # Fuzzy match: allow partial category name
        matched = [c for c in by_cat if cat_name.lower() in c.lower() or c.lower() in cat_name.lower()]
        for cat in matched:
            for rd in active_rows(cat):
                if rd["draft"] and not rd["draft"].startswith("[This factor"):
                    analysis_para(rd["draft"], rd["authority"])
                    wrote_anything = True

    if not wrote_anything:
        para(f"[Attorney to complete analysis for: {', '.join(cat_list) or 'this section'}.]",
             italic=True, size=11)


for sec in primary_sections:
    if not sec["heading"]:
        continue
    letter = section_letters[arg_idx]
    write_argument_section(letter, sec["heading"], sec["categories"])
    arg_idx += 1


# ── AOS Is the Appropriate Mechanism (Part 3) ────────────────────────────────
if aos_include:
    letter = section_letters[arg_idx]
    arg_idx += 1
    heading(f"    {letter}.  {aos_heading}", level=2, space_before=10, space_after=4)

    aos_fn = FN.register(
        "INA §245(a); INA §245(c); 8 CFR §245.1; "
        "USCIS Policy Memorandum PM-602-0199 (May 21, 2026); 1 USCIS-PM E.8"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(6)
    p.add_run(
        "Congress created adjustment of status to permit eligible applicants who were lawfully "
        "admitted to obtain permanent residence without the needless disruption and family "
        "separation that consular processing may require. "
        f"{CLIENT_NAME} was lawfully inspected and admitted into the United States and falls "
        "squarely within the class of aliens Congress authorized to adjust status under "
        "INA §245(a). Consular processing is not required here and, in this case, requiring "
        "the Applicant to depart would [attorney to specify: e.g., trigger a 3- or 10-year bar, "
        "separate the Applicant from dependent family members during an indefinite adjudication, "
        "etc.]. Adjustment of status therefore advances, rather than frustrates, the "
        "family-unification purposes underlying the immigration laws. The Applicant's lawful "
        "entry, immediate-relative petition, and the family dependence described above are "
        "precisely the circumstances adjustment of status was designed to address."
    ).font.size = Pt(12)
    add_fn_ref(p, aos_fn)

    if aos_notes:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p3.paragraph_format.space_before = Pt(3)
        p3.paragraph_format.space_after  = Pt(6)
        p3.paragraph_format.left_indent  = Inches(0.25)
        p3.add_run(aos_notes).font.size  = Pt(12)


# ── Adverse Factors & Rebuttal ────────────────────────────────────────────────
if all_adverse_cats:
    letter = section_letters[arg_idx]
    arg_idx += 1
    heading(f"    {letter}.  {adv_heading}", level=2, space_before=10, space_after=4)

    adv_sub_letters = "abcdefghijk"
    for sub_idx, cat in enumerate(all_adverse_cats):
        display = cat.split(" (")[0]
        heading(f"        {sub_idx + 1}.  {display}", level=3, bold=False, size=12, space_before=6)
        for rd in active_rows(cat):
            if rd["draft"]:
                analysis_para(rd["draft"], rd["authority"])
            if rd["rebuttal"] and rd["rebuttal"] not in ("N/A", ""):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after  = Pt(6)
                p.paragraph_format.left_indent  = Inches(0.25)
                run = p.add_run(rd["rebuttal"])
                run.font.size = Pt(12); run.italic = True
else:
    letter = section_letters[arg_idx]
    arg_idx += 1
    heading(f"    {letter}.  The Record Presents No Substantial Adverse Factors",
            level=2, space_before=10, space_after=4)
    para(
        f"The record does not present any adverse factors that would counsel against a favorable "
        f"exercise of discretion in this case. Accordingly, under Matter of Arai, adjustment "
        f"should ordinarily be granted.",
        italic=False
    )


# ── Balancing ─────────────────────────────────────────────────────────────────
letter = section_letters[arg_idx]
bal_fn = FN.register(
    "Matter of Marin, 16 I&N Dec. 581, 584–85 (BIA 1978); "
    "Matter of Arai, 13 I&N Dec. 494, 496 (BIA 1970); "
    "1 USCIS-PM E.8(C)(3)"
)
heading(f"    {letter}.  The Totality of the Circumstances Warrants a Favorable Exercise of Discretion",
        level=2, space_before=10, space_after=4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(6)
p.add_run(
    f"When the record is viewed in its totality, the equities presented by {CLIENT_NAME} are "
    "substantial and compelling. The Applicant's [attorney: identify two or three strongest equities] "
    "are precisely the kind of social and humane considerations the BIA has recognized as "
    "warranting a favorable exercise of discretion. "
    + (balancing_notes.rstrip(".") + ". " if balancing_notes else "")
    + "Any adverse factors in this record are, by comparison, limited in weight and do not "
    "rise to a level requiring the 'unusual or even outstanding equities' reserved for more "
    "serious cases."
).font.size = Pt(12)
add_fn_ref(p, bal_fn)

# Closing theme sentence ties it all together
if case_theme:
    p_close = doc.add_paragraph()
    p_close.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_close.paragraph_format.space_before = Pt(4)
    p_close.paragraph_format.space_after  = Pt(8)
    p_close.add_run(
        f"This is not a case about an adverse factor. {case_theme} "
        f"A favorable exercise of discretion is both legally supported and consistent with the "
        f"humanitarian purposes Congress built into the adjustment of status framework."
    ).font.size = Pt(12)

hrule()


# ══════════════════════════════════════════════════════════════════════════════
# SECTION IV — CONCLUSION
# ══════════════════════════════════════════════════════════════════════════════
heading("IV.  CONCLUSION", level=1, space_before=16)
para(
    f"For the foregoing reasons, {CLIENT_NAME} respectfully requests that USCIS approve the "
    f"Application to Register Permanent Residence or Adjust Status (Form I-485) and exercise "
    f"its discretion favorably. The Applicant satisfies every statutory requirement under "
    f"INA §245(a), and the totality of the record — including {case_theme.lower() if case_theme else 'the substantial positive equities described above'} "
    f"— demonstrates that a favorable exercise of discretion is both warranted and consistent "
    f"with the intent of the immigration laws. "
    f"{conclusion_notes}"
)

hrule()


# ══════════════════════════════════════════════════════════════════════════════
# FOOTNOTES PAGE
# ══════════════════════════════════════════════════════════════════════════════
doc.add_page_break()
heading("FOOTNOTES", level=1, size=12, space_before=0, space_after=8)

for fn_num, citation in FN.notes:
    p = doc.add_paragraph()
    p.paragraph_format.space_before        = Pt(1)
    p.paragraph_format.space_after         = Pt(4)
    p.paragraph_format.left_indent         = Inches(0.3)
    p.paragraph_format.first_line_indent   = Inches(-0.3)
    ref_run = p.add_run(f"{fn_num}  ")
    ref_run.bold = True
    ref_run.font.size = Pt(10)
    cit_run = p.add_run(citation)
    cit_run.font.size = Pt(10)


# ══════════════════════════════════════════════════════════════════════════════
# SAVE + PATCH settings.xml (zoom attribute)
# ══════════════════════════════════════════════════════════════════════════════
doc.save(OUTPUT_PATH)

_tmp = OUTPUT_PATH + ".tmp"
try:
    with zipfile.ZipFile(OUTPUT_PATH, "r") as zin:
        with zipfile.ZipFile(_tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/settings.xml":
                    data = re.sub(rb"<w:zoom\s*/>", b'<w:zoom w:percent="100"/>', data)
                    data = re.sub(
                        rb"(<w:zoom\b(?![^>]*w:percent)[^>]*/?>)",
                        lambda m: m.group(0).replace(b"<w:zoom", b'<w:zoom w:percent="100"', 1),
                        data,
                    )
                zout.writestr(item, data)
    os.replace(_tmp, OUTPUT_PATH)
except Exception as e:
    print(f"[warn] settings.xml patch failed: {e}")
    if os.path.exists(_tmp):
        os.remove(_tmp)

print(f"Brief written to: {OUTPUT_PATH}")
print(f"Total footnotes:  {len(FN.notes)}")
print(f"Theme sections:   {arg_idx}")
